# Dùng: python3 build_docx.py <file.md> ...   -> tạo .docx cạnh thư mục cha (Tom_tat_bai_bao)
import sys, os, re, json
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
NAVY=RGBColor(0x1F,0x3A,0x5F)
def fld(p,code):
    r=p.add_run(); a=OxmlElement('w:fldChar'); a.set(qn('w:fldCharType'),'begin')
    i=OxmlElement('w:instrText'); i.set(qn('xml:space'),'preserve'); i.text=code
    b=OxmlElement('w:fldChar'); b.set(qn('w:fldCharType'),'end')
    r._r.append(a); r._r.append(i); r._r.append(b)
def shade(cell,hexc):
    tcPr=cell._tc.get_or_add_tcPr(); s=OxmlElement('w:shd'); s.set(qn('w:val'),'clear'); s.set(qn('w:color'),'auto'); s.set(qn('w:fill'),hexc); tcPr.append(s)
def runs(p,text,size=None):
    for part in re.split(r'(\*\*.+?\*\*)',text):
        if part.startswith('**') and part.endswith('**'):
            r=p.add_run(part[2:-2]); r.bold=True
        else: r=p.add_run(part)
        if size: r.font.size=Pt(size)
def build(md,out):
    lines=open(md,encoding='utf-8').read().split('\n')
    doc=Document(); sec=doc.sections[0]; sec.page_width=Cm(21); sec.page_height=Cm(29.7)
    sec.left_margin=sec.right_margin=Cm(2.3); sec.top_margin=Cm(2); sec.bottom_margin=Cm(2)
    st=doc.styles['Normal']; st.font.name='Calibri'; st.font.size=Pt(11.5); st.element.rPr.rFonts.set(qn('w:eastAsia'),'Calibri')
    fp=sec.footer.paragraphs[0]; fp.alignment=WD_ALIGN_PARAGRAPH.CENTER; fp.add_run('Trang '); fld(fp,'PAGE')
    i=0
    while i<len(lines):
        l=lines[i].rstrip()
        if not l: i+=1; continue
        if l.startswith('# '):
            p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(6); r=p.add_run(l[2:]); r.bold=True; r.font.size=Pt(18); r.font.color.rgb=NAVY
        elif l.startswith('## '):
            p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(12); p.paragraph_format.space_after=Pt(3); p.paragraph_format.keep_with_next=True
            r=p.add_run(l[3:]); r.bold=True; r.font.size=Pt(14); r.font.color.rgb=NAVY
        elif l.startswith('### '):
            p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(8); p.paragraph_format.space_after=Pt(2); p.paragraph_format.keep_with_next=True
            r=p.add_run(l[4:]); r.bold=True; r.font.size=Pt(12); r.font.color.rgb=RGBColor(0x33,0x55,0x88)
        elif l.startswith('- '):
            p=doc.add_paragraph(style='List Bullet'); p.paragraph_format.space_after=Pt(2); runs(p,l[2:])
        elif l.startswith('|'):
            rows=[]
            while i<len(lines) and lines[i].startswith('|'):
                cells=[c.strip() for c in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r':?-+:?',c) for c in cells): rows.append(cells)
                i+=1
            nc=max(len(r) for r in rows); t=doc.add_table(rows=len(rows),cols=nc); t.style='Table Grid'
            for ri,r in enumerate(rows):
                for ci in range(nc):
                    c=t.cell(ri,ci); c.text=''; txt=r[ci] if ci<len(r) else ''
                    runs(c.paragraphs[0],txt,10.5)
                    if ri==0:
                        shade(c,'E8EEF5')
                        for rr in c.paragraphs[0].runs: rr.bold=True
            doc.add_paragraph().paragraph_format.space_after=Pt(2)
            continue
        else:
            p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(5); p.paragraph_format.line_spacing=1.15; runs(p,l)
        i+=1
    doc.save(out)
if __name__=='__main__':
    here=os.path.dirname(os.path.abspath(__file__))
    for f in sys.argv[1:]:
        out=os.path.join(os.path.dirname(here),os.path.basename(f)[:-3]+'.docx'); build(f,out); print(out)
